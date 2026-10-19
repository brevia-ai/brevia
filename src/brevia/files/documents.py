"""Document readers: the text of PDF, DOCX and text files, with its structure.

:func:`read_document` reads a file into a :class:`Document`, a sequence of
:class:`Section` values:

- PDF: one section per page with text, with its page number;
- DOCX: one section per heading (paragraph styles ``Title`` and ``Heading N``, or
  styles based on them); paragraphs and tables separated by a blank line, with the
  text before the first heading in a section without title; table rows are lines with
  their cells separated by ``" | "``, so that the values of adjacent cells stay apart;
- text (``.txt``, ``.md``, ``.csv``, ``.json``): one section, read as UTF-8.

The format follows the file extension. Text is cleaned of NUL characters and of runs of
blank lines. The warnings of the PDF parser are not logged: they may quote the file.
Reading is blocking: in an activity use ``ctx.files.read``, which reads in a thread.
Scanned PDF files have no text and are refused: OCR is not supported.
"""

from __future__ import annotations

import logging
import re
from pathlib import Path
from typing import TYPE_CHECKING, Any, Literal

from pydantic import BaseModel, ConfigDict

from brevia.errors import InvalidInput, Unsupported

if TYPE_CHECKING:
    from docx.table import Table
    from docx.text.paragraph import Paragraph

Format = Literal["pdf", "docx", "text"]

FORMATS: dict[str, Format] = {
    ".pdf": "pdf",
    ".docx": "docx",
    ".txt": "text",
    ".md": "text",
    ".csv": "text",
    ".json": "text",
}
"""Readable extensions and their format."""

SEPARATOR = "\n\n"
"""Separator of sections and paragraphs in the text of a document."""

_BLANK_LINES = re.compile(r"\n[ \t]*\n(?:[ \t]*\n)+")
_HEADING_STYLES = re.compile(r"^(Title|Heading \d+)$")


class Section(BaseModel):
    """Part of a document: a page of a PDF, a heading of a DOCX, a whole text file.

    Attributes:
        title: Heading of the section, ``None`` when it has none.
        page: Page number (from 1) for PDF files, otherwise ``None``.
        text: Text of the section, without its title.
    """

    model_config = ConfigDict(frozen=True, extra="forbid")

    title: str | None = None
    page: int | None = None
    text: str


class Document(BaseModel):
    """Text of a file, divided into sections in reading order.

    Attributes:
        name: Name of the file.
        format: ``"pdf"``, ``"docx"`` or ``"text"``.
        sections: Sections in reading order; at least one.
    """

    model_config = ConfigDict(frozen=True, extra="forbid")

    name: str
    format: Format
    sections: tuple[Section, ...]

    @classmethod
    def from_text(cls, text: str, *, name: str = "text") -> Document:
        """Document of one section with ``text``, cleaned as the readers do.

        Raises:
            InvalidInput: The text is empty once cleaned.
        """
        cleaned = clean_text(text)
        if not cleaned:
            raise InvalidInput(
                "The document has no text", {"format": "text", "reason": "no_text"}
            )
        return cls(name=name, format="text", sections=(Section(text=cleaned),))

    @property
    def text(self) -> str:
        """Whole text: the sections with their titles, separated by blank lines."""
        return "\n\n".join(_section_text(section) for section in self.sections)


def read_document(path: str | Path, *, name: str | None = None) -> Document:
    """Read the text of a PDF, DOCX or text file.

    Args:
        path: File to read; its extension gives the format.
        name: Name recorded in the document; the file name by default.

    Raises:
        Unsupported: The extension is not a readable format.
        InvalidInput: The file cannot be read: damaged, encrypted, not UTF-8 text, or
            without any text (for example a scanned PDF).
        FileNotFoundError: The file does not exist.
    """
    path = Path(path)
    check_readable(path.name)
    format_ = FORMATS[path.suffix.lower()]
    if format_ == "pdf":
        sections = _read_pdf(path)
    elif format_ == "docx":
        sections = _read_docx(path)
    else:
        sections = _read_text(path)
    if not sections:
        raise InvalidInput(
            "The document has no text",
            {"format": format_, "reason": "no_text"},
        )
    return Document(name=name or path.name, format=format_, sections=tuple(sections))


def clean_text(text: str) -> str:
    """Text without NUL characters, trailing spaces and runs of blank lines."""
    text = text.replace("\x00", "").replace("\r\n", "\n").replace("\r", "\n")
    text = "\n".join(line.rstrip() for line in text.split("\n"))
    return _BLANK_LINES.sub("\n\n", text).strip()


def _section_text(section: Section) -> str:
    if section.title and section.text:
        return f"{section.title}\n\n{section.text}"
    return section.title or section.text


def _unreadable(format_: Format, reason: str) -> InvalidInput:
    # The exception of the library is not kept: its message may quote the content.
    return InvalidInput(
        "The document cannot be read", {"format": format_, "reason": reason}
    )


def _read_pdf(path: Path) -> list[Section]:
    if not path.exists():
        raise FileNotFoundError(path)
    # Any failure of the parser means a damaged file. The error is raised outside the
    # handler, so that it keeps no reference to the exception of the library, whose
    # message may quote the content.
    damaged = False
    try:
        pages = _pdf_pages(path)
    except Exception:
        damaged = True
    if damaged:
        raise _unreadable("pdf", "damaged")
    if pages is None:
        raise _unreadable("pdf", "encrypted")
    sections = [
        Section(page=number, text=text) for number, text in enumerate(pages, 1) if text
    ]
    if not sections:
        raise InvalidInput(
            "The PDF has no text: OCR of scanned pages is not supported",
            {"format": "pdf", "reason": "no_text", "pages": len(pages)},
        )
    return sections


def _pdf_pages(path: Path) -> list[str] | None:
    """Text of each page, ``None`` for a PDF encrypted with a password."""
    from pypdf import PdfReader

    # pypdf logs its warnings with bytes of the file (an invalid header quotes the
    # first ones): only its critical messages are kept, for every logger under pypdf.
    logging.getLogger("pypdf").setLevel(logging.CRITICAL)

    reader = PdfReader(path)
    if reader.is_encrypted and not reader.decrypt(""):
        return None
    return [clean_text(page.extract_text() or "") for page in reader.pages]


def _read_docx(path: Path) -> list[Section]:
    if not path.exists():
        raise FileNotFoundError(path)
    # As for PDF files: any failure of the parser, raised outside the handler.
    damaged = False
    try:
        sections = _docx_sections(path)
    except Exception:
        damaged = True
    if damaged:
        raise _unreadable("docx", "damaged")
    return sections


def _docx_sections(path: Path) -> list[Section]:
    import docx
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    sections: list[Section] = []
    title: str | None = None
    # Paragraphs and tables are blocks separated by a blank line, as the paragraphs
    # that split() divides a long section at; the rows of a table stay on one line
    # each, in one block.
    blocks: list[str] = []

    def close() -> None:
        text = clean_text(SEPARATOR.join(blocks))
        if title or text:
            sections.append(Section(title=title, text=text))

    for item in docx.Document(str(path)).iter_inner_content():
        if isinstance(item, Paragraph):
            if _is_heading(item) and item.text.strip():
                close()
                title, blocks = clean_text(item.text), []
            else:
                blocks.append(item.text)
        elif isinstance(item, Table):
            blocks.append("\n".join(_table_lines(item)))
    close()
    return sections


def _is_heading(paragraph: Paragraph) -> bool:
    """A paragraph whose style is a heading, or is based on one."""
    style: Any = paragraph.style
    for _ in range(10):  # the chain of base styles is short; guard against loops
        if style is None:
            return False
        if _HEADING_STYLES.match(style.name or ""):
            return True
        style = style.base_style
    return False


def _table_lines(table: Table) -> list[str]:
    lines = []
    # The cell elements are kept, not their ids: a released proxy may reuse an id.
    seen: set[Any] = set()
    for row in table.rows:
        cells: list[str] = []
        for cell in row.cells:
            # A merged cell is repeated in every grid column and row it spans.
            if cell._tc in seen:
                continue
            seen.add(cell._tc)
            cells.append(" ".join(cell.text.split()))
        if cells:
            lines.append(" | ".join(cells))
    return lines


def _read_text(path: Path) -> list[Section]:
    try:
        text = path.read_bytes().decode("utf-8-sig")
    except UnicodeDecodeError:
        raise _unreadable("text", "not_utf8") from None
    text = clean_text(text)
    return [Section(text=text)] if text else []


class Piece(BaseModel):
    """Consecutive text of a document that fits a size, made by :func:`split`.

    Attributes:
        text: Text of the piece: whole sections with their titles, or part of one
            section longer than the size.
        start: Index in ``Document.sections`` of the first section of the piece.
        end: Index after the last section of the piece; ``end - start`` is 1 for a
            part of a long section.
    """

    model_config = ConfigDict(frozen=True, extra="forbid")

    text: str
    start: int
    end: int


def split(document: Document, max_chars: int) -> list[Piece]:
    """Divide ``document`` into pieces of at most ``max_chars`` characters.

    Whole sections are kept together, in order, as long as they fit. A section longer
    than ``max_chars`` is divided at its paragraphs (blank lines), and a paragraph
    still longer at ``max_chars``. Pieces do not overlap; sections without text are
    skipped.

    Raises:
        ValueError: ``max_chars`` is not positive.
    """
    if max_chars <= 0:
        raise ValueError("max_chars must be positive")
    pieces: list[Piece] = []
    texts: list[str] = []
    length = 0
    start = 0

    def close(end: int) -> None:
        nonlocal length
        if texts:
            pieces.append(Piece(text="\n\n".join(texts), start=start, end=end))
            texts.clear()
            length = 0

    for index, section in enumerate(document.sections):
        text = _section_text(section)
        if not text:
            continue
        if len(text) > max_chars:
            close(index)
            pieces.extend(
                Piece(text=part, start=index, end=index + 1)
                for part in _parts(text, max_chars)
            )
            continue
        if texts and length + len(SEPARATOR) + len(text) > max_chars:
            close(index)
        if not texts:
            start = index
        length += len(text) + (len(SEPARATOR) if texts else 0)
        texts.append(text)
    close(len(document.sections))
    return pieces


def check_readable(name: str) -> None:
    """Check that the extension of ``name`` is a readable format.

    Raises:
        Unsupported: The extension is not a readable format.
    """
    suffix = Path(name).suffix.lower()
    if suffix not in FORMATS:
        raise Unsupported(
            "File format not readable",
            {"extension": suffix, "readable": sorted(FORMATS)},
        )


def _parts(text: str, max_chars: int) -> list[str]:
    """Paragraphs of ``text`` grouped within ``max_chars``; longer ones are cut."""
    parts: list[str] = []
    current: list[str] = []
    length = 0
    for paragraph in text.split(SEPARATOR):
        if len(paragraph) > max_chars:
            if current:
                parts.append(SEPARATOR.join(current))
                current, length = [], 0
            parts.extend(
                paragraph[offset : offset + max_chars]
                for offset in range(0, len(paragraph), max_chars)
            )
            continue
        if current and length + len(SEPARATOR) + len(paragraph) > max_chars:
            parts.append(SEPARATOR.join(current))
            current, length = [], 0
        length += len(paragraph) + (len(SEPARATOR) if current else 0)
        current.append(paragraph)
    if current:
        parts.append(SEPARATOR.join(current))
    return parts
