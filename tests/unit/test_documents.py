"""Document readers on synthetic PDF, DOCX and text files written by the tests."""

from __future__ import annotations

import logging
import random
import zipfile
from pathlib import Path

import docx
import pytest
from docx.enum.style import WD_STYLE_TYPE
from pypdf import PdfWriter

from brevia import (
    Document,
    FileSource,
    FileStore,
    InvalidInput,
    Piece,
    RunFiles,
    Section,
    Unsupported,
    read_document,
    split,
)
from brevia.files.documents import clean_text

pytestmark = pytest.mark.unit


def write_pdf(path: Path, pages: list[str]) -> Path:
    """A PDF with one line of Helvetica text per page; an empty string gives a page
    without text."""
    objects = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"",  # pages, written below
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    kids = []
    for text in pages:
        stream = f"BT /F1 12 Tf 72 720 Td ({text}) Tj ET".encode() if text else b""
        objects.append(
            b"<< /Length %d >>\nstream\n%s\nendstream" % (len(stream), stream)
        )
        content = len(objects)
        objects.append(
            b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
            b"/Resources << /Font << /F1 3 0 R >> >> /Contents %d 0 R >>" % content
        )
        kids.append(b"%d 0 R" % len(objects))
    objects[1] = b"<< /Type /Pages /Kids [%s] /Count %d >>" % (
        b" ".join(kids),
        len(kids),
    )
    out = bytearray(b"%PDF-1.4\n")
    offsets = []
    for number, body in enumerate(objects, 1):
        offsets.append(len(out))
        out += b"%d 0 obj\n%s\nendobj\n" % (number, body)
    xref = len(out)
    out += b"xref\n0 %d\n0000000000 65535 f \n" % (len(objects) + 1)
    out += b"".join(b"%010d 00000 n \n" % offset for offset in offsets)
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF\n" % (
        len(objects) + 1,
        xref,
    )
    path.write_bytes(bytes(out))
    return path


def test_pdf_has_one_section_per_page_with_text(tmp_path: Path) -> None:
    path = write_pdf(tmp_path / "report.pdf", ["First page", "", "Third page"])

    document = read_document(path)

    assert document == Document(
        name="report.pdf",
        format="pdf",
        sections=(
            Section(page=1, text="First page"),
            Section(page=3, text="Third page"),
        ),
    )
    assert document.text == "First page\n\nThird page"


def test_pdf_without_text_is_refused(tmp_path: Path) -> None:
    path = write_pdf(tmp_path / "scan.pdf", ["", ""])

    with pytest.raises(InvalidInput) as raised:
        read_document(path)

    assert raised.value.details == {"format": "pdf", "reason": "no_text", "pages": 2}


def test_encrypted_pdf_is_refused(tmp_path: Path) -> None:
    writer = PdfWriter()
    writer.add_blank_page(width=612, height=792)
    writer.encrypt("secret-password")
    path = tmp_path / "locked.pdf"
    with path.open("wb") as file:
        writer.write(file)

    with pytest.raises(InvalidInput) as raised:
        read_document(path)

    assert raised.value.details == {"format": "pdf", "reason": "encrypted"}


@pytest.mark.parametrize("name", ["damaged.pdf", "damaged.docx"])
def test_damaged_file_is_refused_without_its_content(tmp_path: Path, name: str) -> None:
    path = tmp_path / name
    path.write_bytes(b"not a document: private text")

    with pytest.raises(InvalidInput) as raised:
        read_document(path)

    assert raised.value.details["reason"] == "damaged"
    assert "private" not in str(raised.value)
    assert raised.value.__cause__ is None
    assert raised.value.__context__ is None


def test_a_damaged_pdf_logs_nothing_of_its_content(
    tmp_path: Path, caplog: pytest.LogCaptureFixture
) -> None:
    path = tmp_path / "damaged.pdf"
    path.write_bytes(b"SECRET private text of the file")
    caplog.set_level(logging.DEBUG)

    with pytest.raises(InvalidInput):
        read_document(path)

    # pypdf warns about the invalid header quoting its first bytes ("SECRE").
    logged = " ".join(record.getMessage() for record in caplog.records)
    assert "SECRE" not in logged
    assert "private" not in logged


def test_docx_with_a_malformed_part_is_damaged(tmp_path: Path) -> None:
    valid = tmp_path / "valid.docx"
    docx.Document().save(str(valid))
    path = tmp_path / "broken.docx"
    with zipfile.ZipFile(valid) as source, zipfile.ZipFile(path, "w") as target:
        for item in source.infolist():
            content = source.read(item)
            if item.filename == "word/document.xml":
                content = b"<w:document><unclosed private text"
            target.writestr(item, content)

    with pytest.raises(InvalidInput) as raised:
        read_document(path)

    assert raised.value.details == {"format": "docx", "reason": "damaged"}
    assert raised.value.__context__ is None


def test_docx_heading_styles_based_on_headings_and_vertical_merges(
    tmp_path: Path,
) -> None:
    source = docx.Document()
    custom = source.styles.add_style("Section title", WD_STYLE_TYPE.PARAGRAPH)
    custom.base_style = source.styles["Heading 1"]
    source.add_paragraph("Costs", style="Section title")
    table = source.add_table(rows=2, cols=2)
    table.cell(0, 0).merge(table.cell(1, 0)).text = "2025"
    table.cell(0, 1).text = "10"
    table.cell(1, 1).text = "20"
    path = tmp_path / "custom.docx"
    source.save(str(path))

    assert read_document(path).sections == (
        Section(title="Costs", text="2025 | 10\n20"),
    )


def test_docx_has_one_section_per_heading_and_separate_table_cells(
    tmp_path: Path,
) -> None:
    source = docx.Document()
    source.add_paragraph("Before any heading.")
    source.add_heading("Report", level=0)
    source.add_paragraph("Introduction.")
    source.add_heading("Results", level=1)
    source.add_paragraph("Values:")
    table = source.add_table(rows=2, cols=3)
    for cell, text in zip(table.rows[0].cells, ["Year", "Cost", "Value"], strict=True):
        cell.text = text
    for cell, text in zip(table.rows[1].cells, ["2025", "10", "20"], strict=True):
        cell.text = text
    merged = source.add_table(rows=1, cols=3)
    merged.cell(0, 0).merge(merged.cell(0, 1)).text = "Total"
    merged.cell(0, 2).text = "30"
    source.add_heading("Empty", level=2)
    path = tmp_path / "report.docx"
    source.save(str(path))

    document = read_document(path, name="Assembly minutes")

    assert document.name == "Assembly minutes"
    assert document.format == "docx"
    assert document.sections == (
        Section(text="Before any heading."),
        Section(title="Report", text="Introduction."),
        Section(
            title="Results",
            text="Values:\n\nYear | Cost | Value\n2025 | 10 | 20\n\nTotal | 30",
        ),
        Section(title="Empty", text=""),
    )
    assert document.text.endswith("Total | 30\n\nEmpty")


@pytest.mark.parametrize("suffix", [".txt", ".md", ".csv", ".json"])
def test_text_file_is_one_section(tmp_path: Path, suffix: str) -> None:
    path = tmp_path / f"notes{suffix}"
    path.write_bytes("﻿Line one\x00\r\n\r\n\r\n\r\nLine two  \n".encode())

    document = read_document(path)

    assert document.format == "text"
    assert document.sections == (Section(text="Line one\n\nLine two"),)


def test_text_not_in_utf8_is_refused(tmp_path: Path) -> None:
    path = tmp_path / "latin.txt"
    path.write_bytes("caffè".encode("latin-1"))

    with pytest.raises(InvalidInput) as raised:
        read_document(path)

    assert raised.value.details == {"format": "text", "reason": "not_utf8"}


def test_empty_text_file_is_refused(tmp_path: Path) -> None:
    path = tmp_path / "empty.txt"
    path.write_text(" \n\n ")

    with pytest.raises(InvalidInput) as raised:
        read_document(path)

    assert raised.value.details == {"format": "text", "reason": "no_text"}


def test_unknown_extension_is_unsupported(tmp_path: Path) -> None:
    path = tmp_path / "slides.pptx"
    path.write_bytes(b"x")

    with pytest.raises(Unsupported) as raised:
        read_document(path)

    assert raised.value.details["extension"] == ".pptx"


@pytest.mark.parametrize("name", ["missing.pdf", "missing.docx", "missing.txt"])
def test_missing_file_raises_file_not_found(tmp_path: Path, name: str) -> None:
    with pytest.raises(FileNotFoundError):
        read_document(tmp_path / name)


def test_clean_text_removes_nul_and_blank_runs() -> None:
    assert clean_text("a\x00b\n \n\t\n\nc\r\nd ") == "ab\n\nc\nd"


# ---------- documents in runs ----------


async def test_run_files_read_a_source_with_its_name(tmp_path: Path) -> None:
    original = tmp_path / "upload-123.docx"
    source_docx = docx.Document()
    source_docx.add_heading("Summary", level=1)
    source_docx.add_paragraph("Body.")
    source_docx.save(str(original))
    files = RunFiles(FileStore(tmp_path / "files"), "run-1")

    document = await files.read(FileSource(name="minutes.docx", path=original))

    assert document.name == "minutes.docx"
    assert document.sections == (Section(title="Summary", text="Body."),)
    await files.close()
    assert not (tmp_path / "files" / "work" / "run-1").exists()


async def test_run_files_read_follows_the_extension_of_the_source_name(
    tmp_path: Path,
) -> None:
    original = tmp_path / "notes.txt"
    original.write_text("plain")
    files = RunFiles(FileStore(tmp_path / "files"), "run-1")

    with pytest.raises(InvalidInput):
        await files.read(FileSource(name="notes.pdf", path=original))
    await files.close()


def test_document_from_text_is_cleaned() -> None:
    document = Document.from_text("a\x00\n\n\n\nb", name="request")

    assert document == Document(
        name="request", format="text", sections=(Section(text="a\n\nb"),)
    )
    with pytest.raises(InvalidInput):
        Document.from_text(" \n ")


# ---------- pieces ----------


def sections_document(*sections: Section) -> Document:
    return Document(name="doc", format="pdf", sections=sections)


def test_short_document_is_one_piece() -> None:
    document = sections_document(
        Section(page=1, text="one"), Section(page=2, text="two")
    )

    assert split(document, 100) == [Piece(text="one\n\ntwo", start=0, end=2)]


def test_sections_are_grouped_while_they_fit() -> None:
    document = sections_document(
        *(Section(page=n, text=f"page {n} " + "x" * 10) for n in range(1, 6))
    )

    pieces = split(document, 40)

    assert [(p.start, p.end) for p in pieces] == [(0, 2), (2, 4), (4, 5)]
    assert all(len(p.text) <= 40 for p in pieces)
    assert "\n\n".join(p.text for p in pieces) == document.text


def test_titles_stay_with_their_section() -> None:
    document = Document(
        name="doc",
        format="docx",
        sections=(Section(title="Intro", text="a" * 5), Section(title="End", text="b")),
    )

    assert split(document, 12) == [
        Piece(text="Intro\n\naaaaa", start=0, end=1),
        Piece(text="End\n\nb", start=1, end=2),
    ]


def test_a_long_section_is_divided_at_paragraphs_then_at_the_size() -> None:
    long_text = "first paragraph\n\nsecond\n\n" + "z" * 25
    document = sections_document(
        Section(page=1, text="short"),
        Section(page=2, text=long_text),
        Section(page=3, text="after"),
    )

    pieces = split(document, 10)

    assert pieces == [
        Piece(text="short", start=0, end=1),
        Piece(text="first para", start=1, end=2),
        Piece(text="graph", start=1, end=2),
        Piece(text="second", start=1, end=2),
        Piece(text="z" * 10, start=1, end=2),
        Piece(text="z" * 10, start=1, end=2),
        Piece(text="z" * 5, start=1, end=2),
        Piece(text="after", start=2, end=3),
    ]
    assert all(len(p.text) <= 10 for p in pieces)


def test_paragraphs_of_a_long_section_are_grouped_while_they_fit() -> None:
    document = sections_document(Section(page=1, text="aaa\n\nbbb\n\nccc\n\nddd"))

    assert [p.text for p in split(document, 8)] == ["aaa\n\nbbb", "ccc\n\nddd"]


def test_size_must_be_positive() -> None:
    with pytest.raises(ValueError, match="positive"):
        split(sections_document(Section(text="a")), 0)


def test_sections_without_text_are_skipped() -> None:
    document = sections_document(
        Section(page=1, text="a"), Section(text=""), Section(page=3, text="b")
    )

    # The ranges of consecutive pieces stay contiguous: the empty section is inside
    # the range of the first one.
    assert split(document, 3) == [
        Piece(text="a", start=0, end=2),
        Piece(text="b", start=2, end=3),
    ]


def test_pieces_fit_and_keep_the_text_when_no_paragraph_is_cut() -> None:
    generator = random.Random(7)  # noqa: S311 - test data, not security
    for _ in range(300):
        max_chars = generator.randint(20, 80)
        sections = []
        for _ in range(generator.randint(1, 12)):
            paragraphs = [
                "w" * generator.randint(1, max_chars)
                for _ in range(generator.randint(1, 4))
            ]
            title = generator.choice([None, "t" * generator.randint(1, 10)])
            sections.append(Section(title=title, text="\n\n".join(paragraphs)))
        document = Document(name="doc", format="docx", sections=tuple(sections))

        pieces = split(document, max_chars)

        assert all(len(p.text) <= max_chars for p in pieces)
        assert [p.start for p in pieces] == sorted(p.start for p in pieces)
        if all(len(t) <= max_chars for t in document.text.split("\n\n")):
            assert "\n\n".join(p.text for p in pieces) == document.text


async def test_run_files_refuse_an_unreadable_format_before_getting_it(
    tmp_path: Path,
) -> None:
    files = RunFiles(FileStore(tmp_path / "files"), "run-1")

    with pytest.raises(Unsupported):
        await files.read(FileSource(name="deck.pptx", path=tmp_path / "missing.pptx"))
    await files.close()


def test_paragraphs_of_a_docx_are_kept_whole_by_split(tmp_path: Path) -> None:
    # The paragraphs read from a DOCX are those that split() divides a long section
    # at: two paragraphs of six characters, with a size of ten, give two pieces.
    source = docx.Document()
    source.add_paragraph("aaaaaa")
    source.add_paragraph("bbbbbb")
    path = tmp_path / "two.docx"
    source.save(str(path))

    document = read_document(path)

    assert document.sections == (Section(text="aaaaaa\n\nbbbbbb"),)
    assert [piece.text for piece in split(document, 10)] == ["aaaaaa", "bbbbbb"]
