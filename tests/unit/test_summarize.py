"""The ready summarize activity, with the prompts of the package and a fake model."""

from __future__ import annotations

from collections.abc import AsyncIterator
from pathlib import Path
from typing import Any

import docx
import pytest
from conftest import Recorder, simulated
from memory_store import MemoryStore
from pydantic import ValidationError

from brevia import Brevia, FileSource, FileStore, InvalidInput
from brevia.analysis import (
    SUMMARIZE_PROMPTS,
    SummarizeConfig,
    SummarizeRequest,
    Summary,
    summarize,
)
from brevia.config.project import load_project

pytestmark = pytest.mark.unit

MODEL = "openai:test-model"


def project(**defaults: object) -> str:
    lines = "\n".join(
        f"{key} = {value!r}".replace("'", '"') for key, value in defaults.items()
    )
    return f"""
[limits]
max_requests = 20

[models."{MODEL}"]

[activities.summarize]
model = "{MODEL}"
{lines}
"""


async def started(
    tmp_path: Path, rules: str | None = None, **defaults: object
) -> Brevia:
    (tmp_path / "brevia.toml").write_text(project(**defaults), encoding="utf-8")
    prompts: dict[str, str | Path] = dict(SUMMARIZE_PROMPTS)
    if rules is not None:
        (tmp_path / "rules.md").write_text(rules, encoding="utf-8")
        prompts["summarize_rules"] = "rules.md"
    brevia = Brevia(
        load_project(tmp_path / "brevia.toml"),
        project_dir=tmp_path,
        store=MemoryStore(),
        files=FileStore(tmp_path / "files"),
    )
    brevia.register_activity(
        name="summarize",
        version="1",
        function=summarize,
        input_type=SummarizeRequest,
        output_type=Summary,
        config_type=SummarizeConfig,
        prompts=prompts,
        overridable=["strategy", "model.max_tokens"],
        modes=("direct", "job"),
    )
    await brevia.start()
    return brevia


@pytest.fixture
async def app(tmp_path: Path) -> AsyncIterator[Brevia]:
    brevia = await started(tmp_path, max_input_chars=80)
    yield brevia
    await brevia.stop()


def echo(instructions: str, text: str) -> dict[str, Any]:
    return {"summary": f"S[{text[:5]}]"}


def kind(instructions: str) -> str:
    """Name of the prompt of the package that ``instructions`` were rendered from."""
    for name, path in SUMMARIZE_PROMPTS.items():
        if instructions.startswith(path.read_text(encoding="utf-8").split("{")[0]):
            return name
    raise AssertionError(instructions)


async def test_a_short_text_is_summarised_in_one_call(app: Brevia) -> None:
    recorder = Recorder(echo)
    with simulated(recorder.model()):
        result = await app.run("summarize", SummarizeRequest(text="A short text."))

    assert result == Summary(summary="S[A sho]")
    [(instructions, text)] = recorder.calls
    assert text == "A short text."
    assert kind(instructions) == "summarize"
    # Without a language, the summary keeps the language of the document.
    assert "in the language of the document." in instructions.splitlines()[0]
    assert "Do not add information" in instructions


async def test_an_override_changes_one_option_of_the_model(app: Brevia) -> None:
    with simulated(Recorder(echo).model()):
        await app.run("summarize", {"text": "Text."}, {"model": {"max_tokens": 77}})

    store = app._store
    assert isinstance(store, MemoryStore)
    (run,) = store.runs.values()
    # The model of the project is kept: only its option changes.
    assert run.snapshot["config"]["model"]["model"] == MODEL
    assert run.snapshot["config"]["model"]["max_tokens"] == 77


async def test_the_language_of_the_request_prevails(app: Brevia) -> None:
    recorder = Recorder(echo)
    with simulated(recorder.model()):
        await app.run("summarize", SummarizeRequest(text="Text.", language="English"))

    assert "in English." in recorder.calls[0][0]


async def test_the_configured_language_applies_without_one_in_the_request(
    tmp_path: Path,
) -> None:
    brevia = await started(tmp_path, language="Italian")
    recorder = Recorder(echo)
    try:
        with simulated(recorder.model()):
            await brevia.run("summarize", SummarizeRequest(text="Text."))
    finally:
        await brevia.stop()

    assert "in Italian." in recorder.calls[0][0]


async def test_a_long_document_is_summarised_by_pieces_then_combined(
    app: Brevia,
) -> None:
    text = "\n\n".join(f"Paragraph {n} " + "x" * 20 for n in range(4))
    recorder = Recorder(echo)
    with simulated(recorder.model()):
        result = await app.run("summarize", SummarizeRequest(text=text))

    # Paragraphs of 32 characters: two per piece of at most 80, then one combination.
    assert [kind(instructions) for instructions, _ in recorder.calls] == [
        "summarize",
        "summarize",
        "summarize_combine",
    ]
    assert all(len(user) <= 80 for _, user in recorder.calls)
    partial = Summary(summary="S[Parag]").model_dump_json(indent=2)
    assert recorder.calls[2][1] == f"{partial}\n\n{partial}"
    assert result == Summary(summary='S[{\n  "]')


async def test_the_refine_strategy_updates_one_summary(app: Brevia) -> None:
    text = "\n\n".join(f"Paragraph {n} " + "x" * 20 for n in range(3))
    recorder = Recorder(echo)
    with simulated(recorder.model()):
        await app.run(
            "summarize", SummarizeRequest(text=text), overrides={"strategy": "refine"}
        )

    assert [kind(instructions) for instructions, _ in recorder.calls] == [
        "summarize",
        "summarize_refine",
    ]
    assert '"summary": "S[Parag]"' in recorder.calls[1][0]


async def test_a_file_is_read_and_summarised(app: Brevia, tmp_path: Path) -> None:
    path = tmp_path / "minutes.docx"
    document = docx.Document()
    document.add_heading("Minutes", level=1)
    document.add_paragraph("Approved.")
    document.save(str(path))
    recorder = Recorder(echo)
    with simulated(recorder.model()):
        await app.run(
            "summarize",
            SummarizeRequest(file=FileSource(name="minutes.docx", path=path)),
        )

    assert recorder.calls[0][1] == "Minutes\n\nApproved."


async def test_an_empty_text_is_refused_before_any_call(app: Brevia) -> None:
    recorder = Recorder(echo)
    with simulated(recorder.model()), pytest.raises(InvalidInput):
        await app.run("summarize", SummarizeRequest(text=" \n "))

    assert recorder.calls == []


def test_the_request_needs_exactly_one_source(tmp_path: Path) -> None:
    with pytest.raises(ValidationError, match="exactly one"):
        SummarizeRequest()
    with pytest.raises(ValidationError, match="exactly one"):
        SummarizeRequest(
            text="a", file=FileSource(name="a.txt", path=tmp_path / "a.txt")
        )


def test_the_prompts_are_files_of_the_package() -> None:
    assert set(SUMMARIZE_PROMPTS) == {
        "summarize",
        "summarize_combine",
        "summarize_refine",
    }
    assert all(path.is_file() for path in SUMMARIZE_PROMPTS.values())


async def test_a_text_of_exactly_the_size_takes_one_call(app: Brevia) -> None:
    recorder = Recorder(echo)
    with simulated(recorder.model()):
        await app.run("summarize", SummarizeRequest(text="y" * 80))

    assert len(recorder.calls) == 1


async def test_both_sources_are_refused_as_invalid_input(
    app: Brevia, tmp_path: Path
) -> None:
    both = {"text": "a", "file": {"name": "a.txt", "path": str(tmp_path / "a.txt")}}
    with pytest.raises(InvalidInput):
        await app.run("summarize", both)
    with pytest.raises(InvalidInput):
        await app.run("summarize", {})


@pytest.mark.parametrize("strategy", ["auto", "refine"])
async def test_rules_of_the_project_reach_every_call(
    tmp_path: Path, strategy: str
) -> None:
    brevia = await started(
        tmp_path, rules="Only the decisions of the assembly.", max_input_chars=80
    )
    text = "\n\n".join(f"Paragraph {n} " + "x" * 20 for n in range(4))
    recorder = Recorder(echo)
    try:
        with simulated(recorder.model()):
            await brevia.run(
                "summarize",
                SummarizeRequest(text=text),
                overrides={"strategy": strategy},
            )
    finally:
        await brevia.stop()

    assert len(recorder.calls) == (3 if strategy == "auto" else 2)
    assert all(
        instructions.endswith("\n\nOnly the decisions of the assembly.")
        for instructions, _ in recorder.calls
    )


async def test_without_rules_the_prompts_of_the_package_are_used_alone(
    app: Brevia,
) -> None:
    recorder = Recorder(echo)
    with simulated(recorder.model()):
        await app.run("summarize", SummarizeRequest(text="Short."))

    rendered = SUMMARIZE_PROMPTS["summarize"].read_text(encoding="utf-8")
    assert recorder.calls[0][0] == rendered.strip().replace(
        "{language}", "the language of the document"
    )
